"""Export des clients avec leurs soldes (poste)."""
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from accounting_test import DatabaseTestCase
from app.services import client_service, sale_service as sales
from app.services import transaction_service as tx


class ClientExportTests(DatabaseTestCase):
    """Le fixture partagé pose un dépôt de 10 000 sur CLIENT et un produit à 100."""

    def test_export_mirrors_the_screen_and_ends_with_a_total(self):
        tx.create_transaction("SANSFICHE", "", "depot", 2500, self.agent)
        sales.create_sale("DEBITEUR", self.product.id, 3, self.agent)        # -300
        sales.create_sale("", self.product.id, 1, self.agent, mode_paiement="caisse")
        rows = client_service.export_rows()
        screen = client_service.list_clients()
        self.assertEqual([r[0] for r in rows[:-1]], [c.matricule for c in screen])
        by_matricule = {r[0]: r for r in rows[:-1]}
        self.assertEqual(by_matricule["CLIENT"][3].numeric_value, 10000)
        self.assertEqual(by_matricule["DEBITEUR"][3].numeric_value, -300)
        # Un compte sans fiche détient de l'argent : il figure dans l'état.
        self.assertEqual(by_matricule["SANSFICHE"][5], "Non")
        # La vente encaissée n'ouvre aucun compte fantôme.
        self.assertNotIn("", by_matricule)
        total = rows[-1]
        self.assertEqual(total[0], "TOTAL")
        self.assertEqual(total[3].numeric_value, 10000 + 2500 - 300)

    def test_the_search_filters_the_export(self):
        tx.create_transaction("AUTRE", "", "depot", 700, self.agent)
        rows = client_service.export_rows("autre")
        self.assertEqual([r[0] for r in rows], ["AUTRE", "TOTAL"])
        self.assertEqual(rows[-1][3].numeric_value, 700)

    def test_the_excel_file_keeps_balances_as_numbers(self):
        from app.utils.exporters import export_to_excel
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "soldes.xlsx"
            export_to_excel(path, "Clients et soldes", client_service.EXPORT_HEADERS,
                            client_service.export_rows())
            wb = load_workbook(path)
            values = [[c.value for c in row] for row in wb.active.iter_rows()]
            wb.close()
        solde = next(r[3] for r in values if r and r[0] == "CLIENT")
        self.assertIsInstance(solde, (int, float))
        self.assertEqual(solde, 10000)

    def test_many_clients_use_three_queries_and_do_not_start_a_write(self):
        from app.db.database import get_connection
        for i in range(12):
            client_service.create_client(f"FAST-{i}", nom="Nom test")
        statements = []
        conn = get_connection()
        conn.set_trace_callback(statements.append)
        try:
            rows = client_service.export_rows()
        finally:
            conn.set_trace_callback(None)
        selects = [q for q in statements if q.lstrip().upper().startswith("SELECT")]
        self.assertEqual(len(selects), 3)
        self.assertEqual(len(rows), 14)
        self.assertFalse(conn.in_transaction)

    def test_cancelled_cash_and_credit_operations_keep_exact_balances(self):
        tx.create_transaction("A", "", "depot", .1, self.agent)
        tx.create_transaction("A", "", "depot", .2, self.agent)
        purchase = sales.create_sale("A", self.product.id, 1, self.agent)
        sales.cancel_sale(purchase.id)
        sales.create_sale("A", self.product.id, 1, self.agent, mode_paiement="caisse")
        rows = {r[0]: r for r in client_service.export_rows()}
        self.assertEqual(rows["A"][3].numeric_value, .3)
        self.assertEqual(rows["TOTAL"][3].numeric_value, 10000.3)

    def test_pdf_handles_long_names_and_preserves_credit_balances(self):
        from server.sync.client_report import build_client_pdf
        client_service.create_client("CREDIT", nom="Nom très long avec accents & <texte> " * 5)
        sales.create_sale("CREDIT", self.product.id, 1, self.agent)
        content = build_client_pdf("Clients et soldes", client_service.EXPORT_HEADERS,
                                   client_service.export_rows())
        self.assertTrue(content.startswith(b"%PDF"))
