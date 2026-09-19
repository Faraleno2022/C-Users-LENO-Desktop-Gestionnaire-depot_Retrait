"""Export des clients avec leurs soldes (Console Web)."""
import io
import uuid

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from sync.models import Client, Product, Sale, Transaction
from web.tests import StockAndBalanceTests
from web.views import _client_balances, _matricule_balance


def tx(matricule, kind, montant, deleted=False):
    return Transaction.objects.create(
        uuid=str(uuid.uuid4()), matricule=matricule, type=kind, montant=montant,
        solde_apres=0, created_at="2026-09-01 10:00:00", deleted=deleted)


def sale(matricule, montant, mode="compte", deleted=False):
    return Sale.objects.create(
        uuid=str(uuid.uuid4()), matricule=matricule, product_id=1, product_uuid="p",
        product_nom="Riz", quantite=1, prix_unitaire=montant, montant_total=montant,
        solde_apres=0, mode_paiement=mode, created_at="2026-09-01 11:00:00", deleted=deleted)


def fiche(matricule, nom="", telephone=""):
    return Client.objects.create(
        uuid=str(uuid.uuid4()), matricule=matricule, nom=nom, telephone=telephone,
        created_at="now", updated_at="now")


class ClientBalanceExportTests(TestCase):
    def setUp(self):
        StockAndBalanceTests.setUp(self)
        # Le fixture partagé pose un dépôt de 10 000 sur MAT-1.

    make_product = StockAndBalanceTests.make_product

    def export(self, fmt="xlsx", **params):
        return self.client.get(reverse("web:clients"), {"export": fmt, **params})

    def sheet_rows(self, response):
        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active
        return [[c.value for c in row] for row in ws.iter_rows(min_row=4)]

    # --- Exactitude ----------------------------------------------------------

    def test_every_balance_matches_the_per_account_calculation(self):
        """Le calcul groupé doit égaler, au franc près, celui de toute l'application."""
        fiche("A1", "Awa")
        tx("A1", "depot", 40000); tx("A1", "retrait", 7500.5); sale("A1", 12000.25)
        tx("B2", "depot", 0.1); tx("B2", "depot", 0.2)            # piège flottant
        sale("C3", 5000)                                          # crédit sans dépôt
        tx("D4", "depot", 9000, deleted=True)                     # annulé
        for row in _client_balances():
            self.assertEqual(row["solde"], _matricule_balance(row["matricule"]),
                             row["matricule"])
        soldes = {r["matricule"]: r["solde"] for r in _client_balances()}
        self.assertEqual(soldes["A1"], 20499.25)
        self.assertEqual(soldes["B2"], 0.3)       # et non 0.30000000000000004
        self.assertEqual(soldes["C3"], -5000)

    def test_cash_sales_and_deleted_operations_are_left_out(self):
        fiche("E5")
        tx("E5", "depot", 3000)
        sale("E5", 1000, mode="caisse")           # encaissée : aucun compte touché
        sale("E5", 500, deleted=True)             # annulée
        row = next(r for r in _client_balances() if r["matricule"] == "E5")
        self.assertEqual(row["solde"], 3000)
        # Le décompte suit celui du poste : toute vente non annulée compte comme
        # une opération, même encaissée ; seule la vente annulée disparaît.
        self.assertEqual(row["n_ops"], 2)

    def test_an_account_without_a_fiche_is_included(self):
        """Ces comptes détiennent de l'argent : les omettre fausserait l'état."""
        tx("ORPHELIN", "depot", 25000)
        row = next(r for r in _client_balances() if r["matricule"] == "ORPHELIN")
        self.assertFalse(row["fiche"])
        self.assertEqual(row["solde"], 25000)

    # --- Fichier Excel -------------------------------------------------------

    def test_excel_export_has_numeric_balances_and_a_correct_total(self):
        fiche("A1", "Awa", "622000000")
        tx("A1", "depot", 40000); sale("A1", 12000)
        tx("ORPHELIN", "depot", 25000)
        response = self.export("xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertIn("clients_soldes_", response["Content-Disposition"])
        rows = self.sheet_rows(response)
        by_matricule = {r[0]: r for r in rows}
        self.assertEqual(by_matricule["A1"][1:4], ["Awa", "622000000", 28000])
        self.assertEqual(by_matricule["A1"][5], "Oui")
        self.assertEqual(by_matricule["ORPHELIN"][5], "Non")
        # Le solde est un nombre, pas un texte : Excel peut le recalculer.
        self.assertIsInstance(by_matricule["A1"][3], (int, float))
        total = rows[-1]
        self.assertEqual(total[0], "TOTAL")
        comptes = [r for r in rows if r[0] != "TOTAL"]
        self.assertEqual(total[3], sum(r[3] for r in comptes))
        self.assertEqual(total[3], 10000 + 28000 + 25000)   # MAT-1 + A1 + ORPHELIN

    def test_a_credit_balance_is_exported_as_negative(self):
        sale("DEBITEUR", 7000)
        rows = {r[0]: r for r in self.sheet_rows(self.export("xlsx"))}
        self.assertEqual(rows["DEBITEUR"][3], -7000)

    def test_the_search_filters_the_export(self):
        fiche("A1", "Awa"); tx("A1", "depot", 100)
        fiche("B2", "Bintou"); tx("B2", "depot", 200)
        rows = self.sheet_rows(self.export("xlsx", q="bint"))
        self.assertEqual([r[0] for r in rows], ["B2", "TOTAL"])
        self.assertEqual(rows[-1][3], 200)

    def test_the_export_is_not_capped_like_the_page(self):
        """La page s'arrête à 500 comptes ; l'état des soldes, lui, est complet."""
        Client.objects.bulk_create([
            Client(uuid=str(uuid.uuid4()), matricule=f"M{i:04d}", nom="",
                   created_at="now", updated_at="now") for i in range(520)])
        rows = self.sheet_rows(self.export("xlsx"))
        self.assertEqual(len(rows) - 1, 521)     # 520 fiches + MAT-1, hors total

    def test_a_formula_in_a_client_name_stays_plain_text(self):
        fiche("F6", "=HYPERLINK(\"http://x\",\"clic\")")
        tx("F6", "depot", 100)
        wb = load_workbook(io.BytesIO(self.export("xlsx").content))
        cell = next(row[1] for row in wb.active.iter_rows(min_row=4) if row[0].value == "F6")
        self.assertNotEqual(cell.data_type, "f")

    def test_pdf_export(self):
        response = self.export("pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_the_page_offers_both_exports_carrying_the_search(self):
        page = self.client.get(reverse("web:clients"), {"q": "Awa Diallo"}).content.decode()
        self.assertIn("export=xlsx", page)
        self.assertIn("export=pdf", page)
        self.assertIn("q=Awa%20Diallo", page)

    def test_export_requires_login(self):
        self.client.logout()
        self.assertEqual(self.export().status_code, 302)

    def test_empty_export_has_zero_total(self):
        Transaction.objects.all().delete()
        response = self.export("xlsx")
        self.assertEqual(self.sheet_rows(response), [["TOTAL", "0 compte(s)", None, 0, 0, None]])
        self.assertTrue(self.export("pdf").content.startswith(b"%PDF"))

    def test_client_count_does_not_multiply_queries(self):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext
        for i in range(12):
            fiche(f"FAST-{i}")
        with CaptureQueriesContext(connection) as queries:
            rows = _client_balances()
        selects = [q for q in queries if q["sql"].lstrip().upper().startswith("SELECT")]
        self.assertEqual(len(selects), 3)
        self.assertEqual(len(rows), 13)
