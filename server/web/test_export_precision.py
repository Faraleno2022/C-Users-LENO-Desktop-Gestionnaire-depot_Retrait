import io
from django.test import TestCase
from openpyxl import load_workbook
from sync.models import StockMovement
from web import test_pages
from web.reports import build_excel, _fmt_num, _fmt_money


class ExportPrecisionTests(TestCase):
    setUp = test_pages.PageTests.setUp
    def test_exact_numbers_and_text_identifiers(self):
        data = build_excel('Précision', ['Quantité', 'Montant', 'Matricule'],
            [[_fmt_num(1000007397), _fmt_money(123456789.25), '0012'],
             [_fmt_num(999999947), _fmt_money(0), '=1+1']])
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        self.assertEqual((ws['A4'].value, ws['A4'].data_type), (1000007397, 'n'))
        self.assertEqual(ws['A5'].value, 999999947)
        self.assertEqual(ws['B4'].value, 123456789.25)
        self.assertIn('GNF', ws['B4'].number_format)
        self.assertEqual((ws['C4'].value, ws['C4'].data_type), ('0012', 's'))
        self.assertEqual(ws['C5'].data_type, 's')
        self.assertEqual(str(_fmt_num(1000007397)), '1000007397')
        wb.close()

    def test_movement_summary_and_detail_keep_exact_values(self):
        p = self.product
        p.stock_initial = 1000000000
        p.stock_initial_source = 'creation'
        p.quantite_stock = 24707
        p.created_at = '2026-09-01 08:00:00'
        p.save()
        for uuid, kind, qty, after, hour in [('entry', 'entree', 32104, 1000032104, 9),
                                           ('exit', 'sortie', 1000007397, 24707, 10)]:
            StockMovement.objects.create(uuid=uuid, product_uuid=p.uuid, product_id=p.id,
                product_nom=p.nom, type=kind, quantite=qty, stock_apres=after,
                created_at=f'2026-09-01 {hour:02}:00:00')
        from django.urls import reverse
        response = self.client.get(reverse('web:stock_movements'), {'export':'xlsx'})
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(io.BytesIO(response.content)); ws = wb.active
        self.assertEqual(ws['D4'].value, 1000007397)
        self.assertEqual(ws['B4'].value + ws['C4'].value - ws['D4'].value, ws['E4'].value)
        self.assertEqual(ws['D7'].value, 1000007397)
        self.assertEqual(ws['E4'].value, ws['F4'].value)
        wb.close()

    def test_inventory_count_cells_stay_blank_and_stock_is_numeric(self):
        from django.urls import reverse
        response = self.client.get(reverse('web:inventory'), {'export':'xlsx'})
        wb = load_workbook(io.BytesIO(response.content)); ws = wb.active
        self.assertEqual(ws['D4'].data_type, 'n')
        self.assertIsNone(ws['E4'].value)
        wb.close()
        response = self.client.get(reverse('web:inventory'))
        self.assertContains(response, 'À compter')
        self.assertContains(response, 'Aucun comptage saisi')
