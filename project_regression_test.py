"""Régressions projet : uniquement bases et fichiers temporaires."""
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
import sqlite3
import unittest

from accounting_test import DatabaseTestCase
from app.db import database as db
from app.services import backup_service as backup, client_service as clients
from app.services import product_service as products, sync_service as sync
from app.services import transaction_service as tx, user_service as users
from app.utils.exporters import export_to_excel
from openpyxl import load_workbook


class ProjectRegressionTests(DatabaseTestCase):
    def test_excel_preserves_large_quantities_money_and_identifiers(self):
        from app.utils.report_values import report_number, report_money
        path = Path(self.tmp.name) / 'precision.xlsx'
        export_to_excel(path, 'Stock', ['Référence', 'Quantité', 'Montant'],
            [['0012', report_number(1000007397), report_money(123456789.25)]])
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        sheet = workbook.active
        self.assertEqual((sheet['A4'].value, sheet['A4'].data_type), ('0012', 's'))
        self.assertEqual((sheet['B4'].value, sheet['B4'].data_type), (1000007397, 'n'))
        self.assertEqual(sheet['C4'].value, 123456789.25)
        self.assertIn('GNF', sheet['C4'].number_format)

    def setup_backup(self):
        folder = Path(self.tmp.name) / "backups"
        folder.mkdir(exist_ok=True)
        for name, value in (("DB_PATH", db.DB_PATH), ("BACKUP_DIR", folder),
                            ("ensure_directories", lambda: None)):
            p = patch.object(backup, name, value)
            p.start()
            self.addCleanup(p.stop)

    def test_backups_same_second_have_distinct_files(self):
        self.setup_backup()
        with patch.object(backup, "datetime") as clock:
            clock.now.return_value = datetime(2026, 1, 1)
            first = backup.create_backup()
            second = backup.create_backup()
        self.assertNotEqual(first, second)
        self.assertTrue(first.exists() and second.exists())

    def test_invalid_restore_keeps_current_database(self):
        self.setup_backup()
        invalid = Path(self.tmp.name) / "invalid.db"
        invalid.write_text("not sqlite")
        with self.assertRaises(backup.BackupError):
            backup.restore_backup(invalid)
        self.assertEqual(tx.get_matricule_balance("CLIENT"), 10000)

    def test_restore_updates_existing_connection_without_wal_replay(self):
        self.setup_backup()
        saved = backup.create_backup()
        tx.create_transaction("CLIENT", "", "retrait", 1000, self.agent)
        observer = sqlite3.connect(str(db.DB_PATH))
        self.addCleanup(observer.close)
        observer.execute("PRAGMA journal_mode=WAL")
        backup.restore_backup(saved)
        self.assertEqual(tx.get_matricule_balance("CLIENT"), 10000)
        self.assertTrue(any(b["note"] == "Avant restauration" for b in backup.list_backups()))
        self.assertEqual(observer.execute("SELECT COUNT(*) FROM transactions").fetchone()[0], 1)

    def test_used_client_matricule_cannot_detach_balance(self):
        client = clients.create_client("CLIENT")
        with self.assertRaises(clients.ClientError):
            clients.update_client(client.id, matricule="AUTRE")
        self.assertEqual(clients.get_client(client.id).matricule, "CLIENT")

    def test_deleting_user_preserves_history_and_syncs_deactivation(self):
        another = users.create_user("other", "password", "Autre", "caissier")
        tx.create_transaction("OTHER", "", "depot", 1, another)
        users.delete_user(another.id)
        self.assertFalse(users.get_user(another.id).actif)
        self.assertEqual(tx.get_matricule_balance("OTHER"), 1)

    def test_last_superadmin_cannot_be_demoted(self):
        users.update_user(self.agent.id, role="super_admin")
        with self.assertRaises(users.UserServiceError):
            users.update_user(self.agent.id, role="caissier")

    def test_push_ack_does_not_hide_changes_made_during_http(self):
        records = sync._collect_pending("products", sync.PUSH_TABLES["products"], 100)
        products.adjust_stock(self.product.id, "entree", 1)
        # New API acknowledges precisely the snapshot sent, not just its UUID.
        sync._mark_synced("products", records)
        row = db.get_connection().execute("SELECT sync_status FROM products WHERE id=?", (self.product.id,)).fetchone()
        self.assertEqual(row["sync_status"], "pending")

    def test_pull_missing_agent_is_retained_without_inventing_identity(self):
        record = dict(db.get_connection().execute("SELECT * FROM stock_movements LIMIT 1").fetchone())
        record.update(uuid="remote-movement", agent_uuid="", agent_id=None, quantite=2)
        outcome = sync._merge_pulled_record("stock_movements", sync.PULL_TABLES["stock_movements"], record)
        self.assertEqual(outcome, "inserted")
        row = db.get_connection().execute("SELECT agent_id FROM stock_movements WHERE uuid='remote-movement'").fetchone()
        self.assertIsNone(row["agent_id"])

    def test_pull_conflict_does_not_advance_cursor(self):
        class Response:
            status_code = 200
            def json(self):
                return {"records": [{"uuid": "unresolved"}], "next_since": "2026-01-01T00:00:00Z", "has_more": False}
        with patch.dict(sync.PULL_TABLES, {"products": sync.PULL_TABLES["products"]}, clear=True), patch.object(
            sync.settings_service, "get_sync_config", return_value={"server_url": "https://example.test", "device_token": "test"}
        ), patch.object(sync.requests, "get", return_value=Response()), patch.object(
            sync, "_merge_pulled_record", return_value="conflict"
        ):
            result = sync.pull_all()
        self.assertEqual(result["products"]["conflict"], 1)
        self.assertFalse(sync.settings_service.get_setting("pull_since_products"))

    def test_excel_handles_sheet_punctuation_and_literal_formulas(self):
        file = Path(self.tmp.name) / "export.xlsx"
        export_to_excel(file, "Dépôts / Retraits : [clients]", ["Nom"], [["=1+1"]])
        wb = load_workbook(file)
        try:
            self.assertEqual(wb.active["A4"].value, "=1+1")
            self.assertEqual(wb.active["A4"].data_type, "s")
        finally:
            wb.close()

    def test_legacy_agent_migration_preserves_records_indexes_and_sequence(self):
        from app.db.schema import SCHEMA_STATEMENTS
        conn = sqlite3.connect(":memory:")
        conn.row_factory = sqlite3.Row
        self.addCleanup(conn.close)
        for statement in SCHEMA_STATEMENTS:
            conn.execute(statement.replace("agent_id INTEGER,", "agent_id INTEGER NOT NULL,"))
        agent = dict(db.get_connection().execute("SELECT * FROM users WHERE id=?", (self.agent.id,)).fetchone())
        columns = list(agent)
        conn.execute("INSERT INTO users (" + ",".join(columns) + ") VALUES (" + ",".join("?" for _ in columns) + ")", list(agent.values()))
        record = dict(db.get_connection().execute("SELECT * FROM transactions LIMIT 1").fetchone())
        columns = list(record)
        conn.execute("INSERT INTO transactions (" + ",".join(columns) + ") VALUES (" + ",".join("?" for _ in columns) + ")", list(record.values()))
        conn.execute("UPDATE sqlite_sequence SET seq=99 WHERE name='transactions'")
        conn.commit()
        db._allow_missing_agents(conn)
        db._allow_missing_agents(conn)
        self.assertEqual(dict(conn.execute("SELECT * FROM transactions").fetchone()), record)
        self.assertEqual(conn.execute("SELECT seq FROM sqlite_sequence WHERE name='transactions'").fetchone()[0], 99)
        self.assertIsNotNone(conn.execute("SELECT name FROM sqlite_master WHERE name='idx_tx_matricule'").fetchone())
        self.assertEqual(conn.execute("PRAGMA foreign_key_check").fetchall(), [])
        record.update(id=100, uuid="without-agent", agent_id=None)
        conn.execute("INSERT INTO transactions (" + ",".join(columns) + ") VALUES (" + ",".join("?" for _ in columns) + ")", list(record.values()))

    def test_pdf_treats_company_and_client_markup_as_text(self):
        from app.utils.exporters import export_to_pdf
        target = Path(self.tmp.name) / "report.pdf"
        export_to_pdf(target, "Client <broken> & fils", ["Nom"], [["A & B"]],
                      subtitle="<unclosed> société")
        self.assertTrue(target.read_bytes().startswith(b"%PDF"))

    def test_database_export_includes_recent_wal_transactions(self):
        self.setup_backup()
        tx.create_transaction("CLIENT", "", "retrait", 321, self.agent)
        target = Path(self.tmp.name) / "export.db"
        backup.export_database(target)
        with sqlite3.connect(target) as exported:
            self.assertEqual(exported.execute("SELECT COUNT(*) FROM transactions").fetchone()[0], 2)
            self.assertEqual(exported.execute("PRAGMA integrity_check").fetchone()[0], "ok")
        exported.close()
        with self.assertRaises(backup.BackupError):
            backup.export_database(db.DB_PATH)

    def test_console_tolerates_malformed_configuration(self):
        import json
        from server import console_web
        path = Path(self.tmp.name) / "render_sync.json"
        for data in ([], {"enabled": "false", "interval_seconds": "oops", "url": [], "token": 1}):
            path.write_text(json.dumps(data))
            cfg = console_web._read_render_config(Path(self.tmp.name))
            self.assertIs(cfg["enabled"], False)
            self.assertEqual(cfg["interval_seconds"], 5)


if __name__ == "__main__":
    unittest.main()
